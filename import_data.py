import openpyxl
import os
from datetime import datetime
from database import init_db, insert_employees_batch, insert_transactions_batch, replace_all_data

EXCEL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "advance.xlsm")


def read_excel_data(excel_path):
    """Read employees (Adv_Summary) and transactions (Data) from the source workbook."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"File not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, read_only=True, keep_vba=False, data_only=True)

    employees = []
    ws_summary = wb["Adv_Summary"]
    for row in ws_summary.iter_rows(min_row=6, values_only=True):
        code = row[1]
        name = row[2]
        designation = row[3] if row[3] else ""
        if code is not None and name is not None and isinstance(code, (int, float)):
            employees.append((int(code), str(name).strip(), str(designation).strip()))

    transactions = []
    ws_data = wb["Data"]
    for row in ws_data.iter_rows(min_row=6, values_only=True):
        trans_date = row[2]
        emp_name = row[3]
        if trans_date is None or emp_name is None or not isinstance(trans_date, datetime):
            continue

        description = row[4] if row[4] else ""
        try:
            advance = float(row[5]) if row[5] else 0
            deduction = float(row[6]) if row[6] else 0
            balance = float(row[7]) if row[7] else 0
            running_total = float(row[8]) if row[8] else 0
        except (ValueError, TypeError):
            continue

        transactions.append((
            str(emp_name).strip(),
            trans_date.strftime("%Y-%m-%d"),
            str(description).strip(),
            advance,
            deduction,
            balance,
            running_total,
        ))

    wb.close()
    return {"employees": employees, "transactions": transactions}


def import_from_excel(excel_path, mode="replace"):
    """Sync the database from the Excel workbook.

    mode="replace" wipes employees + transactions and imports everything fresh
    (the Excel file is treated as the source of truth).
    mode="append" only adds employees/transactions that do not yet exist.
    """
    data = read_excel_data(excel_path)
    init_db()

    if mode == "replace":
        replace_all_data(data["employees"], data["transactions"])
    else:
        insert_employees_batch(data["employees"])
        batch_size = 500
        for i in range(0, len(data["transactions"]), batch_size):
            insert_transactions_batch(data["transactions"][i:i + batch_size])

    return {
        "mode": mode,
        "employees": len(data["employees"]),
        "transactions": len(data["transactions"]),
    }


if __name__ == "__main__":
    result = import_from_excel(EXCEL_FILE, mode="replace")
    print(f"Sync complete ({result['mode']} mode):")
    print(f"  Employees imported: {result['employees']}")
    print(f"  Transactions imported: {result['transactions']}")